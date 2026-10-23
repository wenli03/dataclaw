"""文档解析技能：PDF / DOCX / TXT / MD / XLSX → 纯文本。

优先真实解析（pymupdf / python-docx / openpyxl）；解析结果写入会话黑板，
供合同审查等后续工具复用，避免重复读盘。
"""
from __future__ import annotations

import re
from pathlib import Path

from .. import run_query
from ..types import ToolContext, ToolResult, ToolSpec
from . import STATE_DOC, json_safe

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx", ".html", ".htm", ".json"}
MAX_CHARS = 120_000

BUILTIN_CONTRACT = """技术开发与运维服务合同（演示样例）

甲方：华东智造集团有限公司
乙方：汉得灵舟数字科技有限公司
合同编号：HD-DC-2026-0715

第一条 服务内容
乙方向甲方提供数据中台建设与智能报表开发服务，具体范围以附件《工作说明书》为准。乙方应按照双方确认的项目计划完成交付，并配合甲方完成验收。

第二条 付款方式
甲方在合同签订后 90 日内向乙方支付合同总价 100% 的款项。乙方应在收款后开具发票。项目实施过程中的差旅、硬件采购等费用由乙方自行承担。

第三条 违约责任
乙方逾期交付的，每逾期一日按合同总额的 1% 向甲方支付违约金；甲方逾期付款的，每逾期一日按未付款项的 0.1‰ 支付违约金。因乙方原因导致项目失败的，乙方应退还全部已收款项并赔偿甲方损失。

第四条 争议管辖
因本合同引起的或与本合同有关的任何争议，双方应友好协商解决；协商不成的，任何一方均可向甲方所在地人民法院提起诉讼。

第五条 合同期限与续约
本合同有效期自 2026 年 7 月 15 日起至 2027 年 7 月 14 日止。合同期满前 30 日内，如甲方未书面提出异议，本合同自动续约一年，续约条款与本合同一致。

第六条 责任承担
乙方对因本项目产生的任何第三方索赔、数据泄露及业务中断承担无限连带责任，该责任不因合同终止而免除。

第七条 知识产权
本项目形成的软件著作权、数据模型、文档及其他成果的知识产权归属，由双方另行协商确定。在权属确定前，甲方有权自行使用相关成果。

第八条 保密义务
双方对在合作过程中知悉的商业秘密负保密义务，保密期限为合同终止后 1 年。违反保密义务的一方应赔偿对方因此遭受的直接损失。

第九条 合同解除
甲方有权在提前 7 日书面通知后单方解除本合同，且无需承担违约责任；乙方不得单方解除本合同。合同解除后，甲方已支付的费用不予退还。

第十条 验收标准
项目交付后由甲方组织验收，验收是否通过以甲方单方书面确认为准。甲方未在 30 日内提出书面意见的，视为验收通过。

第十一条 不可抗力
因不可抗力导致合同无法履行的，双方互不承担责任。不可抗力包括但不限于自然灾害、战争、政府行为、疫情、市场价格波动、供应链中断及乙方供应商违约等情形。

第十二条 其他
本合同未尽事宜，由双方另行协商。合同附件与本合同具有同等法律效力。本合同一式肆份，双方各执贰份，自双方盖章之日起生效。

（以下无正文）
"""


def extract_text(path: Path | str, name: str = "") -> tuple[str, str]:
    """返回 (text, mime_or_type)；失败抛异常由调用方转 ToolResult。"""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        try:
            import pymupdf  # type: ignore

            with pymupdf.open(str(p)) as doc:
                text = "\n".join(page.get_text("text") for page in doc)
            return text, "pdf"
        except ImportError:
            import fitz  # type: ignore

            with fitz.open(str(p)) as doc:
                text = "\n".join(page.get_text("text") for page in doc)
            return text, "pdf"
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(p))
        parts = [para.text for para in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts), "docx"
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(str(p), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"# Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c) for c in row]
                if any(cells):
                    lines.append(" | ".join(cells))
        wb.close()
        return "\n".join(lines), "xlsx"
    if suffix in (".txt", ".md", ".csv", ".html", ".htm", ".json"):
        raw = p.read_bytes()
        for encoding in ("utf-8", "utf-8-sig", "gb18030"):
            try:
                return raw.decode(encoding), suffix.lstrip(".")
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="replace"), suffix.lstrip(".")
    if suffix == ".doc":
        raise ValueError("暂不支持 .doc 旧格式，请另存为 .docx 或 PDF 后重试")
    raise ValueError(f"不支持的文件类型：{suffix or name}")


def _file_meta(ctx: ToolContext, file_id: str) -> dict | None:
    from ...models import FileModel

    record = ctx.db.get(FileModel, file_id)
    if record is None:
        return None
    return {"fileId": record.id, "name": record.name, "path": record.path, "mime": record.mime}


def _parse_contract_sample(ctx: ToolContext) -> dict:
    """从 contract_sample 选「风险条款最多」的一份合同，并保留 risk_hint 供规则引擎融合。"""
    ds = ctx.datasource()
    if ds is None:
        return {}
    sql = (
        "SELECT contract_name, clause_no, title, content, category, risk_hint "
        "FROM contract_sample ORDER BY contract_name, clause_no"
    )
    result = run_query(ctx.db, ctx.user, getattr(ds, "id", ""), sql, limit=300)
    if not result.get("ok") or not result.get("rows"):
        return {}
    grouped: dict[str, list[dict]] = {}
    for row in result["rows"]:
        grouped.setdefault(str(row.get("contract_name") or "样例合同"), []).append(row)

    def risky_count(rows: list[dict]) -> int:
        total = 0
        for row in rows:
            hint = str(row.get("risk_hint") or "")
            if "风险" in hint and "有利" not in hint:
                total += 1
        return total

    best_name, best_rows = max(grouped.items(), key=lambda kv: (risky_count(kv[1]), len(kv[1])))
    lines = [f"{best_name}", ""]
    hints: list[dict] = []
    for row in sorted(best_rows, key=lambda r: str(r.get("clause_no") or "")):
        clause_no = str(row.get("clause_no") or "")
        title = str(row.get("title") or "")
        content = str(row.get("content") or "")
        lines.append(f"{clause_no} {title}".strip())
        lines.append(content)
        lines.append("")
        hints.append(
            {
                "clause": clause_no,
                "title": title,
                "category": str(row.get("category") or ""),
                "hint": str(row.get("risk_hint") or ""),
                "content": content,
            }
        )
    return {"text": "\n".join(lines), "name": best_name, "hints": hints, "riskyCount": risky_count(best_rows)}


async def parse_text_handler(params: dict, ctx: ToolContext) -> ToolResult:
    file_id = str(params.get("fileId") or params.get("file_id") or "").strip()
    contract_name = str(params.get("contractName") or params.get("contract_name") or "").strip()
    source = "attachment"
    name = ""
    mime = ""
    path = ""

    if not file_id:
        # 从会话附件里兜底
        for att in ctx.attachments or []:
            if att.get("fileId"):
                file_id = att["fileId"]
                break
    if file_id:
        meta = _file_meta(ctx, file_id)
        if meta is None:
            return ToolResult(ok=False, error=f"文件不存在或无权访问：{file_id}")
        name = meta["name"]
        path = meta["path"]
        mime = meta["mime"]
    else:
        sample = _parse_contract_sample(ctx)
        if sample and sample.get("riskyCount", 0) >= 6:
            text = sample["text"]
            source = "contract_sample"
            name = contract_name or sample["name"]
            hints = sample.get("hints") or []
        else:
            text, source = BUILTIN_CONTRACT, "builtin-sample"
            name = contract_name or "技术开发与运维服务合同（演示样例）"
            hints = []
        ctx.state[STATE_DOC] = {
            "docName": name,
            "source": source,
            "text": text,
            "riskHints": hints,
            "paragraphs": [p for p in text.split("\n") if p.strip()],
            "chars": len(text),
            "fileId": "",
            "path": "",
        }
        return ToolResult(
            ok=True,
            summary=f"已解析{('数据库样例合同' if source == 'contract_sample' else '内置样例合同')}《{name}》，共 {len(text)} 字",
            output=json_safe({"docName": name, "source": source, "chars": len(text), "preview": text[:600]}),
        )

    if not path or not Path(path).exists():
        return ToolResult(ok=False, error=f"文件路径不存在：{path}")
    try:
        text, parsed_type = extract_text(path, name)
    except Exception as exc:
        return ToolResult(ok=False, error=f"文档解析失败：{exc}")
    text = (text or "").strip()
    if not text:
        return ToolResult(ok=False, error="文档未解析出文本（可能是扫描件，需 OCR）")
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS]
    paragraphs = [p.strip() for p in re.split(r"\n{1,}", text) if p.strip()]
    ctx.state[STATE_DOC] = {
        "docName": name,
        "source": source,
        "text": text,
        "paragraphs": paragraphs,
        "chars": len(text),
        "fileId": file_id,
        "path": path,
        "mime": mime,
        "parsedType": parsed_type,
    }
    output = json_safe(
        {
            "docName": name,
            "fileId": file_id,
            "source": source,
            "parsedType": parsed_type,
            "chars": len(text),
            "paragraphs": len(paragraphs),
            "preview": text[:800],
            "text": text,
        }
    )
    return ToolResult(ok=True, summary=f"已解析《{name}》{len(text)} 字 / {len(paragraphs)} 段", output=output)


async def list_files_handler(params: dict, ctx: ToolContext) -> ToolResult:
    from ...models import FileModel

    rows = (
        ctx.db.query(FileModel)
        .filter(FileModel.uploaded_by == getattr(ctx.user, "id", ""))
        .order_by(FileModel.created_at.desc())
        .limit(50)
        .all()
    )
    items = [{"fileId": r.id, "name": r.name, "size": r.size, "mime": r.mime, "kind": r.kind} for r in rows]
    return ToolResult(ok=True, summary=f"我的文件共 {len(items)} 个", output={"items": items})


async def save_note_handler(params: dict, ctx: ToolContext) -> ToolResult:
    title = str(params.get("title") or "文档解析笔记")
    content = str(params.get("content") or "").strip()
    if not content:
        doc = ctx.state.get(STATE_DOC) or {}
        content = (doc.get("text") or "")[:3000]
    if not content:
        return ToolResult(ok=False, error="没有可保存的内容")
    return ToolResult(
        ok=True,
        summary=f"已生成笔记《{title}》",
        output={"title": title, "chars": len(content)},
        artifact={
            "kind": "report",
            "title": title,
            "payload": {"type": "report", "markdown": f"# {title}\n\n{content}", "title": title},
            "text": content,
        },
    )


TOOLS: list[ToolSpec] = [
    ToolSpec(
        name="doc.list_files",
        title="列出我的文件",
        description="列出当前用户上传的文件（供文档解析选择）。",
        params={"type": "object", "properties": {}},
        handler=list_files_handler,
        risk="low",
        skill="doc-parse",
    ),
    ToolSpec(
        name="doc.parse_text",
        title="解析文档为文本",
        description="解析 PDF/DOCX/TXT/MD/XLSX 为纯文本；无附件时可读取演示样例合同。",
        params={
            "type": "object",
            "properties": {
                "fileId": {"type": "string", "description": "上传文件 ID"},
                "contractName": {"type": "string", "description": "可选：演示合同名"},
            },
        },
        handler=parse_text_handler,
        risk="low",
        skill="doc-parse",
    ),
    ToolSpec(
        name="doc.save_note",
        title="保存解析笔记",
        description="把解析结果或指定内容保存为可下载的笔记产物。",
        params={
            "type": "object",
            "properties": {
                "title": {"type": "string"},
                "content": {"type": "string"},
            },
        },
        handler=save_note_handler,
        risk="low",
        skill="doc-parse",
    ),
]
