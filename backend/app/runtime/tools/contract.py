"""合同审查技能：规则引擎 + 带修订痕迹的 docx 导出。

规则 12 条覆盖：付款账期、预付比例、违约金对等、管辖法院、自动续约、
无限连带、知识产权归属、保密期限、单方解除、验收标准、不可抗力、发票/争议解决缺失。

docx 生成策略：优先写入 OOXML 修订痕迹（w:ins / w:del，作者 DataClaw），
同时在正文保留红色加粗的「修订建议」段落（可读性兜底）；生成后用 python-docx
回读 + zip 校验，确保 Word 能正常打开。
"""
from __future__ import annotations

import re
import time
import zipfile
from datetime import datetime
from pathlib import Path

from ...core.config import ARTIFACT_DIR
from ..types import ToolContext, ToolResult, ToolSpec
from . import STATE_CONTRACT, STATE_DOC, STATE_KNOWLEDGE, json_safe

RISK_ORDER = {"high": 0, "medium": 1, "low": 2}
AUTHOR = "DataClaw"

CLAUSE_RE = re.compile(
    r"(?m)^\s*(第[一二三四五六七八九十百零〇\d]+条|\d+(?:\.\d+){1,3}|\d+[、.．])\s*([^\n]{0,60})"
)


def split_clauses(text: str) -> list[dict]:
    text = (text or "").strip()
    if not text:
        return []
    matches = list(CLAUSE_RE.finditer(text))
    clauses: list[dict] = []
    if matches:
        for idx, match in enumerate(matches):
            start = match.start()
            end = matches[idx + 1].start() if idx + 1 < len(matches) else len(text)
            raw = text[start:end].strip()
            header = match.group(1)
            title = (match.group(2) or "").strip()
            body = raw[len(header) :].strip()
            clauses.append({"clause": header, "title": title, "text": raw, "body": body})
        return clauses
    # 无明确条款编号：按句号切块，每 3 句一组
    sentences = [s.strip() for s in re.split(r"(?<=。)", text) if s.strip()]
    for idx in range(0, len(sentences), 3):
        chunk = "".join(sentences[idx : idx + 3])
        clauses.append({"clause": f"第{idx // 3 + 1}段", "title": "", "text": chunk, "body": chunk})
    return clauses


def _find(clauses: list[dict], *keywords: str) -> list[dict]:
    return [c for c in clauses if any(k in (c.get("text") or "") for k in keywords)]


def _pct_values(text: str) -> list[float]:
    values = []
    for match in re.finditer(r"(\d+(?:\.\d+)?)\s*[%％]", text):
        try:
            values.append(float(match.group(1)))
        except ValueError:
            continue
    for match in re.finditer(r"(\d+(?:\.\d+)?)\s*‰", text):
        try:
            values.append(float(match.group(1)) / 10.0)
        except ValueError:
            continue
    return values


def _risk(clause: dict, title: str, risk: str, issue: str, suggestion: str) -> dict:
    original = (clause.get("text") or "").strip()
    return {
        "clause": clause.get("clause") or "全文",
        "title": title,
        "risk": risk,
        "issue": issue,
        "suggestion": suggestion,
        "original": original[:600],
    }


# --------------------------------------------------------------------------- #
# 规则引擎
# --------------------------------------------------------------------------- #


def analyze_rules(text: str, clauses: list[dict]) -> list[dict]:
    risks: list[dict] = []
    full = text or ""

    # 1. 付款账期 > 60 天
    for clause in _find(clauses, "付款", "支付", "账期", "结算"):
        body = clause.get("text") or ""
        days = [int(d) for d in re.findall(r"(\d{1,3})\s*(?:个)?(?:日|天)", body)]
        if days and max(days) > 60:
            risks.append(
                _risk(
                    clause,
                    "付款账期过长",
                    "high",
                    f"约定 {max(days)} 天账期，超过公司标准（≤60 天），回款风险与资金占用显著上升。",
                    f"建议将账期调整为 60 天内，或改为「30 天付 60% + 验收后 30 天付 40%」的分期方案。",
                )
            )
            break

    # 2. 预付比例过低 / 无预付
    for clause in _find(clauses, "付款", "预付", "首付"):
        body = clause.get("text") or ""
        if ("预付" not in body and "首付" not in body) and ("100%" in body or "全额" in body):
            risks.append(
                _risk(
                    clause,
                    "无预付款安排",
                    "medium",
                    "约定验收/完成后一次性付款，未设置预付款，乙方需自行垫资。",
                    "建议增加「合同签订后 10 个工作日内支付不低于 30% 预付款」条款。",
                )
            )
            break

    # 3. 违约金比例不对等
    for clause in _find(clauses, "违约", "违约金", "赔偿"):
        values = _pct_values(clause.get("text") or "")
        if len(values) >= 2:
            low, high = min(values), max(values)
            if low > 0 and high / low >= 5:
                risks.append(
                    _risk(
                        clause,
                        "违约金比例不对等",
                        "high",
                        f"双方违约金差距达 {high / low:.0f} 倍（{low}% vs {high}%），存在显失公平与谈判风险。",
                        "建议双方违约金比例对齐（如均为合同总额的 0.05%/日，累计不超过 10%），并设置封顶。",
                    )
                )
                break

    # 4. 管辖法院非我方所在地
    for clause in _find(clauses, "管辖", "法院", "诉讼", "仲裁"):
        body = clause.get("text") or ""
        if "甲方所在地" in body and "乙方所在地" not in body:
            risks.append(
                _risk(
                    clause,
                    "管辖法院不利",
                    "medium",
                    "争议由甲方所在地法院管辖，我方异地应诉成本高。",
                    "建议改为「由被告所在地/合同履行地人民法院管辖」，或约定双方认可的仲裁机构。",
                )
            )
            break

    # 5. 自动续约
    for clause in _find(clauses, "续约", "自动续", "顺延"):
        body = clause.get("text") or ""
        if "自动" in body or ("未" in body and "异议" in body):
            risks.append(
                _risk(
                    clause,
                    "自动续约条款",
                    "medium",
                    "期满未书面异议即自动续约一年，可能形成非预期长期义务。",
                    "建议改为「需双方书面确认后续约」，或至少提前 60 日双方确认机制。",
                )
            )
            break

    # 6. 无限连带责任
    for clause in _find(clauses, "责任", "连带", "赔偿", "索赔"):
        body = clause.get("text") or ""
        if "无限" in body and ("连带" in body or "责任" in body):
            risks.append(
                _risk(
                    clause,
                    "无限连带责任",
                    "high",
                    "约定无限连带责任且不因合同终止免除，责任敞口不可控。",
                    "建议将赔偿责任限定为「直接损失」，并设置不超过合同总额 100% 的责任上限；排除间接损失。",
                )
            )
            break

    # 7. 知识产权归属模糊
    for clause in _find(clauses, "知识产权", "著作权", "成果归属", "专利"):
        body = clause.get("text") or ""
        if any(k in body for k in ("另行协商", "另行约定", "未约定", "另行确定")):
            risks.append(
                _risk(
                    clause,
                    "知识产权归属模糊",
                    "high",
                    "成果权属约定为「另行协商」，交付后易产生权属争议，影响产品复用。",
                    "建议明确：通用组件/平台能力归乙方，定制交付物归甲方；并授予甲方永久使用许可。",
                )
            )
            break

    # 8. 保密期限 < 3 年
    for clause in _find(clauses, "保密"):
        body = clause.get("text") or ""
        years = re.findall(r"(\d+)\s*年", body)
        if years and min(int(y) for y in years) < 3:
            risks.append(
                _risk(
                    clause,
                    "保密期限过短",
                    "medium",
                    f"保密期限仅 {min(int(y) for y in years)} 年，低于公司标准（≥3 年），商业秘密保护不足。",
                    "建议保密期限延长至合同终止后 3-5 年，并对核心数据设置更长期限。",
                )
            )
            break

    # 9. 单方解除权不对等
    for clause in _find(clauses, "解除", "终止"):
        body = clause.get("text") or ""
        if "甲方" in body and "单方解除" in body and ("乙方不得" in body or "乙方无权" in body):
            risks.append(
                _risk(
                    clause,
                    "单方解除权不对等",
                    "high",
                    "甲方可单方解除且无需担责，乙方无对等权利，已投入成本可能无法收回。",
                    "建议增加对等解除权，并约定解除后的费用结算与已完成工作量补偿。",
                )
            )
            break

    # 10. 验收标准不明确
    for clause in _find(clauses, "验收"):
        body = clause.get("text") or ""
        if any(k in body for k in ("单方确认", "以甲方", "视为验收通过", "自行决定")):
            risks.append(
                _risk(
                    clause,
                    "验收标准不明确",
                    "medium",
                    "验收以甲方单方书面确认为准、逾期视为通过，标准主观且不可控。",
                    "建议附件列明可量化验收清单与测试用例，明确默认验收的条件与时限。",
                )
            )
            break

    # 11. 不可抗力范围过宽
    for clause in _find(clauses, "不可抗力"):
        body = clause.get("text") or ""
        if any(k in body for k in ("市场价格", "供应链中断", "供应商违约", "政策变化")):
            risks.append(
                _risk(
                    clause,
                    "不可抗力范围过宽",
                    "low",
                    "将市场价格波动、供应链中断等商业风险纳入不可抗力，可能被滥用免责。",
                    "建议将不可抗力限定为不能预见、不能避免且不能克服的客观事件，商业风险除外。",
                )
            )
            break

    # 12a. 发票/税率条款缺失（文档级）
    if "发票" not in full or not re.search(r"\d+\s*%", full):
        anchor = {"clause": "全文", "text": ""}
        risks.append(
            _risk(
                anchor,
                "发票与税率条款缺失",
                "low",
                "全文未明确发票类型（增值税专用/普通）与税率，影响进项抵扣与结算。",
                "建议补充「乙方开具 6% 增值税专用发票，甲方在收到发票后 15 日内付款」。",
            )
        )

    # 12b. 争议解决方式缺失
    if not re.search(r"仲裁|诉讼|法院|管辖", full):
        anchor = {"clause": "全文", "text": ""}
        risks.append(
            _risk(
                anchor,
                "争议解决条款缺失",
                "medium",
                "未约定仲裁或诉讼及管辖地，发生争议时确定管辖成本高。",
                "建议明确争议解决方式、适用法律与管辖法院/仲裁机构。",
            )
        )

    risks.sort(key=lambda r: (RISK_ORDER.get(r["risk"], 9), r["clause"]))
    return risks


# --------------------------------------------------------------------------- #
# risk_hint 融合 + 类别评估（对 seed 合同与措辞不标准的合同兜底）
# --------------------------------------------------------------------------- #

CATEGORY_SUGGESTIONS = {
    "付款": "建议将账期压缩至 60 天内，或改为「30 天付 60% + 验收后 30 天付 40%」，并明确可量化验收标准。",
    "违约": "建议双方违约金比例对齐（如均为 0.05%/日，累计不超过合同总额 10%），并设置责任上限、排除间接损失。",
    "管辖": "建议改为「被告所在地/合同履行地人民法院管辖」，或约定双方认可的仲裁机构。",
    "续约": "建议改为「需双方书面确认后续约」，并提前 60 日由双方确认。",
    "责任": "建议将赔偿责任限定为直接损失，并设置不超过合同总额 100% 的责任上限。",
    "知识产权": "建议明确：通用组件/平台能力归乙方，定制交付物归甲方，并授予甲方永久使用许可；不得包含乙方既有技术。",
    "保密": "建议保密期限延长至合同终止后 3-5 年，并对核心数据设置更长期限。",
    "解除": "建议增加对等解除权，并约定解除后的费用结算与已完成工作量补偿。",
    "验收": "建议附件列明可量化验收清单与测试用例，明确默认验收条件与时限。",
    "其他": "建议逐条复核该条款的适用条件与责任边界，必要时补充限制性表述。",
    "服务": "建议补充《服务水平协议》与工作说明书附件，明确交付物清单与验收口径。",
}

RISK_CATEGORY = {
    "付款账期过长": "付款",
    "无预付款安排": "付款",
    "发票与税率条款缺失": "付款",
    "验收标准不明确": "付款",
    "违约金比例不对等": "违约",
    "违约责任不对等": "违约",
    "无限连带责任": "责任",
    "管辖法院不利": "管辖",
    "争议解决条款缺失": "管辖",
    "自动续约条款": "续约",
    "知识产权归属模糊": "知识产权",
    "知识产权归属不利": "知识产权",
    "保密期限过短": "保密",
    "单方解除权不对等": "解除",
    "不可抗力范围过宽": "其他",
    "服务范围不明确": "服务",
    "条款约定不完整": "其他",
}

CATEGORY_KEYWORDS = {
    "付款": ("付款", "支付", "账期", "发票", "验收"),
    "违约": ("违约", "赔偿", "逾期"),
    "责任": ("责任", "连带", "上限"),
    "知识产权": ("知识产权", "著作权", "专利", "成果"),
    "保密": ("保密", "商业秘密"),
    "管辖": ("管辖", "法院", "仲裁", "诉讼"),
    "续约": ("续约", "续期", "顺延"),
    "解除": ("解除", "终止"),
}


def risk_category(title: str) -> str:
    if title in RISK_CATEGORY:
        return RISK_CATEGORY[title]
    for key, category in RISK_CATEGORY.items():
        if key in title:
            return category
    return "其他"


def _hint_level(hint: str) -> str:
    if "高风险" in hint:
        return "high"
    if "中风险" in hint:
        return "medium"
    if "低风险" in hint:
        return "low"
    return "low"


def _is_risky_hint(hint: str) -> bool:
    if not hint or "有利" in hint or "清晰" in hint or "合规" in hint:
        return False
    return ("风险" in hint) or any(k in hint for k in ("建议", "缺少", "未", "需", "应"))


def _standard_title(category: str, content: str, hint: str) -> str:
    # 标题判定以条款正文为主，hint 只做辅助，避免同一句被拆成两个标题
    text = content or ""
    if category == "付款":
        if "账期" in text or re.search(r"\d+\s*(?:日|天)", text):
            return "付款账期过长"
        if "验收" in text or "满意" in text:
            return "验收标准不明确"
        return "付款条件不利"
    if category == "违约":
        if any(k in text for k in ("不对等", "100 倍", "100倍", "高于")):
            return "违约金比例不对等"
        return "违约责任不对等"
    return {
        "管辖": "管辖法院不利",
        "续约": "自动续约条款",
        "知识产权": "知识产权归属不利",
        "保密": "保密期限过短",
        "解除": "单方解除权不对等",
        "责任": "无限连带责任",
        "服务": "服务范围不明确",
        "其他": "条款约定不完整",
    }.get(category, "条款约定不完整")


def merge_hint_risks(risks: list[dict], hints: list[dict]) -> list[dict]:
    """把 contract_sample.risk_hint 作为「人工标注的规则命中」融合进风险清单。"""
    out = list(risks)
    by_title: dict[str, dict] = {r["title"]: r for r in out}
    for hint in hints or []:
        text = str(hint.get("hint") or "")
        if not _is_risky_hint(text):
            continue
        category = str(hint.get("category") or "").strip() or risk_category(_standard_title("其他", hint.get("content") or "", text))
        title = _standard_title(category, str(hint.get("content") or ""), text)
        level = _hint_level(text)
        issue = re.sub(r"^(高|中|低)风险[：:]\s*", "", text).strip() or text
        clause_text = str(hint.get("content") or "")
        clause = {"clause": hint.get("clause") or "相关条款", "text": clause_text}
        suggestion = CATEGORY_SUGGESTIONS.get(category) or CATEGORY_SUGGESTIONS["其他"]
        existing = by_title.get(title)
        if existing is not None and existing["clause"] == clause["clause"]:
            if RISK_ORDER.get(level, 9) < RISK_ORDER.get(existing["risk"], 9):
                existing["risk"] = level
            if len(issue) > len(existing.get("issue") or ""):
                existing["issue"] = issue
            continue
        entry = _risk(clause, title, level, issue, suggestion)
        out.append(entry)
        by_title.setdefault(title, entry)
    out.sort(key=lambda r: (RISK_ORDER.get(r["risk"], 9), r["clause"]))
    return out


def category_assessments(risks: list[dict], text: str) -> list[dict]:
    """逐类给出结论（命中→风险项；未命中→未见明显问题 / 建议补充）。"""
    out: list[dict] = []
    for category, keywords in CATEGORY_KEYWORDS.items():
        related = [r for r in risks if risk_category(r["title"]) == category]
        present = any(k in text for k in keywords)
        if related:
            titles = "、".join(dict.fromkeys(r["title"] for r in related))
            out.append(
                {
                    "category": category,
                    "status": "risk",
                    "count": len(related),
                    "conclusion": f"发现 {len(related)} 项风险：{titles}",
                }
            )
        else:
            out.append(
                {
                    "category": category,
                    "status": "clear" if present else "absent",
                    "count": 0,
                    "conclusion": "该类别未见明显问题" if present else "合同中未见该类别条款，建议补充",
                }
            )
    return out


# --------------------------------------------------------------------------- #
# docx 生成
# --------------------------------------------------------------------------- #


def _set_cn_font(run, name: str = "微软雅黑", size=None, color=None, bold=None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _add_tracked(paragraph, text: str, kind: str, author: str, date_iso: str, rev_id: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tag = "w:ins" if kind == "ins" else "w:del"
    wrapper = OxmlElement(tag)
    wrapper.set(qn("w:id"), str(rev_id))
    wrapper.set(qn("w:author"), author)
    wrapper.set(qn("w:date"), date_iso)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79" if kind == "ins" else "C00000")
    rpr.append(color)
    bold = OxmlElement("w:b")
    rpr.append(bold)
    run.append(rpr)
    node = OxmlElement("w:t" if kind == "ins" else "w:delText")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    wrapper.append(run)
    paragraph._p.append(wrapper)


def build_docx(doc_name: str, text: str, risks: list[dict], stats: dict, path: Path) -> dict:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    try:
        from docx.oxml.ns import qn

        style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    except Exception:
        pass

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_cn_font(title.add_run(f"《{doc_name}》审查报告（修订版）"), size=18, bold=True, color="1F2937")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_cn_font(
        meta.add_run(
            f"审查对象：{doc_name}｜审查人：{AUTHOR}（数字员工）｜生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
        ),
        size=9,
        color="6B7280",
    )

    doc.add_heading("一、风险概览", level=1)
    summary = doc.add_paragraph()
    _set_cn_font(
        summary.add_run(
            f"共识别风险 {stats.get('risks', 0)} 项：高风险 {stats.get('high', 0)} 项、"
            f"中风险 {stats.get('medium', 0)} 项、低风险 {stats.get('low', 0)} 项。"
        ),
        size=10.5,
    )
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["序号", "条款", "风险", "问题", "修订建议"]
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        _set_cn_font(cell.paragraphs[0].add_run(header), size=9, bold=True, color="FFFFFF")
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "1F4E79")
        cell._tc.get_or_add_tcPr().append(shd)
    for idx, risk in enumerate(risks[:20], 1):
        row = table.add_row().cells
        values = [str(idx), risk["clause"], {"high": "高", "medium": "中", "low": "低"}.get(risk["risk"], risk["risk"]), risk["issue"], risk["suggestion"]]
        for cell, value in zip(row, values):
            cell.text = ""
            _set_cn_font(cell.paragraphs[0].add_run(str(value)), size=9)

    doc.add_heading("二、原文与修订痕迹", level=1)
    note = doc.add_paragraph()
    _set_cn_font(
        note.add_run("说明：蓝色下划线为 DataClaw 建议插入的内容，红色删除线为建议删除的原文；下方红色段落为可读版修订意见。"),
        size=9,
        color="6B7280",
    )
    tracked_ok = True
    date_iso = datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ")
    rev = 1
    try:
        for risk in risks:
            heading = doc.add_paragraph()
            _set_cn_font(
                heading.add_run(f"{risk['clause']} {risk['title']}（{risk['risk'].upper()}）"),
                size=11,
                bold=True,
                color="B91C1C" if risk["risk"] == "high" else "92400E",
            )
            if risk.get("original"):
                para_del = doc.add_paragraph()
                _set_cn_font(para_del.add_run("原文："), size=9, color="6B7280")
                _add_tracked(para_del, risk["original"][:400], "del", AUTHOR, date_iso, rev)
                rev += 1
            para_ins = doc.add_paragraph()
            _set_cn_font(para_ins.add_run("修订："), size=9, color="6B7280")
            _add_tracked(para_ins, risk["suggestion"], "ins", AUTHOR, date_iso, rev)
            rev += 1
            # 可读性兜底：红色加粗修订意见
            fallback = doc.add_paragraph()
            _set_cn_font(
                fallback.add_run(f"【修订建议｜{AUTHOR}｜{datetime.now().strftime('%Y-%m-%d')}】{risk['suggestion']}"),
                size=10,
                bold=True,
                color="C00000",
            )
    except Exception:
        tracked_ok = False

    if not tracked_ok:
        # 极端兜底：纯红色段落
        for risk in risks:
            para = doc.add_paragraph()
            _set_cn_font(
                para.add_run(f"【修订建议】{risk['clause']}：{risk['suggestion']}（原文：{risk.get('original', '')[:200]}）"),
                size=10,
                bold=True,
                color="C00000",
            )

    doc.add_heading("三、审查依据与免责说明", level=1)
    basis = doc.add_paragraph()
    _set_cn_font(
        basis.add_run(
            "本报告依据公司《标准合同审查规则》自动生成，共匹配 12 类风险规则；"
            "审查结果供法务与业务复核参考，不替代正式法律意见。所有修订均保留作者与时间，"
            "可在 Word「审阅」中逐条接受或拒绝。"
        ),
        size=9,
        color="4B5563",
    )
    doc.save(str(path))
    return {"trackedChanges": tracked_ok, "revisions": rev - 1}


def verify_docx(path: Path) -> dict:
    from docx import Document

    doc = Document(str(path))
    paragraphs = len(doc.paragraphs)
    has_ins = has_del = False
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
        has_ins = "<w:ins " in xml
        has_del = "<w:del " in xml
    return {"paragraphs": paragraphs, "hasInsertions": has_ins, "hasDeletions": has_del, "bytes": path.stat().st_size}


# --------------------------------------------------------------------------- #
# 工具
# --------------------------------------------------------------------------- #


async def parse_clauses_handler(params: dict, ctx: ToolContext) -> ToolResult:
    text = str(params.get("text") or (ctx.state.get(STATE_DOC) or {}).get("text") or "")
    if not text:
        return ToolResult(ok=False, error="没有待解析的合同文本，请先执行 doc.parse_text")
    clauses = split_clauses(text)
    ctx.state["contract_clauses"] = clauses
    return ToolResult(
        ok=True,
        summary=f"识别到 {len(clauses)} 个条款",
        output={
            "count": len(clauses),
            "clauses": [{"clause": c["clause"], "title": c["title"], "chars": len(c["text"])} for c in clauses],
        },
    )


async def risk_scan_handler(params: dict, ctx: ToolContext) -> ToolResult:
    doc = ctx.state.get(STATE_DOC) or {}
    text = str(params.get("text") or doc.get("text") or "")
    if not text.strip():
        return ToolResult(ok=False, error="没有合同文本，请先上传合同或执行 doc.parse_text")
    clauses = ctx.state.get("contract_clauses") or split_clauses(text)
    hints = doc.get("riskHints") or []
    risks = merge_hint_risks(analyze_rules(text, clauses), hints)
    assessments = category_assessments(risks, text)
    display_name = Path(str(doc.get("docName") or params.get("contractName") or "合同")).stem or "合同"
    stats = {
        "risks": len(risks),
        "high": sum(1 for r in risks if r["risk"] == "high"),
        "medium": sum(1 for r in risks if r["risk"] == "medium"),
        "low": sum(1 for r in risks if r["risk"] == "low"),
    }
    hits = ctx.state.get(STATE_KNOWLEDGE) or []
    ctx.state[STATE_CONTRACT] = {
        "docName": display_name,
        "docText": text,
        "risks": risks,
        "stats": stats,
        "clauses": len(clauses),
        "categoryAssessments": assessments,
        "riskHintCount": len([h for h in hints if _is_risky_hint(str(h.get("hint") or ""))]),
        "knowledgeHits": hits[:3],
    }
    output = json_safe(
        {
            "docName": display_name,
            "stats": stats,
            "clauses": len(clauses),
            "risks": risks,
            "redlines": risks,
            "rulesMatched": len({r["title"] for r in risks}),
            "categoryAssessments": assessments,
            "knowledgeUsed": [h.get("docName") for h in hits[:3]],
        }
    )
    return ToolResult(
        ok=True,
        summary=f"合同审查完成：《{display_name}》{stats['risks']} 项风险"
        f"（高 {stats['high']} / 中 {stats['medium']} / 低 {stats['low']}），覆盖 {len(clauses)} 个条款",
        output=output,
    )


async def redline_docx_handler(params: dict, ctx: ToolContext) -> ToolResult:
    data = ctx.state.get(STATE_CONTRACT)
    if not data:
        return ToolResult(ok=False, error="请先完成合同风险扫描（contract.risk_scan）")
    if not (ctx.approved.get(ctx.step_id) or ctx.auto_approve):
        return ToolResult(ok=False, error="该操作需要人工审批：将导出含企业条款的修订版文件")
    risks = data.get("risks") or []
    doc_name = str(data.get("docName") or "合同")
    stats = data.get("stats") or {}
    ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
    stem = Path(doc_name).stem or doc_name
    safe = re.sub(r"[\\/:*?\"<>|\s]+", "-", stem)[:60] or "合同"
    file_name = f"合同审查-{safe}-修订版-{int(time.time())}.docx"
    path = Path(ARTIFACT_DIR) / file_name
    build_info = build_docx(doc_name, data.get("docText") or "", risks, stats, path)
    verify = verify_docx(path)
    download_url = f"/api/artifacts/{{id}}/download"  # runner 保存后回填真实 id
    redlines = [
        {
            "clause": r["clause"],
            "risk": r["risk"],
            "title": r["title"],
            "issue": r["issue"],
            "suggestion": r["suggestion"],
            "original": r.get("original", ""),
        }
        for r in risks
    ]
    payload = {
        "type": "docx",
        "title": f"《{stem}》审查报告（修订版）",
        "fileName": file_name,
        "filePath": str(path),
        "downloadUrl": download_url,
        "stats": stats,
        "redlines": redlines,
        "trackedChanges": bool(build_info.get("trackedChanges")),
        "revisions": build_info.get("revisions", 0),
        "verified": verify,
        "author": AUTHOR,
    }
    ctx.state["contract_docx"] = payload
    return ToolResult(
        ok=True,
        summary=f"已生成带修订痕迹的 docx（{stats.get('risks', 0)} 项风险，"
        f"{'含' if verify.get('hasInsertions') else '无'} OOXML 修订标记，{verify.get('bytes', 0) // 1024}KB）",
        output=json_safe(
            {
                "fileName": file_name,
                "filePath": str(path),
                "stats": stats,
                "redlineCount": len(redlines),
                "trackedChanges": payload["trackedChanges"],
                "verified": verify,
            }
        ),
        artifact={
            "kind": "docx",
            "title": payload["title"],
            "payload": json_safe(payload),
            "text": (data.get("docText") or "")[:2000],
            "file_name": file_name,
            "file_path": str(path),
        },
    )


TOOLS: list[ToolSpec] = [
    ToolSpec(
        name="contract.parse_clauses",
        title="解析合同条款",
        description="把合同文本切分为条款（条款号 + 标题 + 正文），供规则匹配。",
        params={"type": "object", "properties": {"text": {"type": "string", "description": "可选：直接传入合同全文"}}},
        handler=parse_clauses_handler,
        risk="low",
        skill="contract-review",
    ),
    ToolSpec(
        name="contract.risk_scan",
        title="合同风险规则扫描",
        description="按 12 类标准合同审查规则匹配风险条款，输出 clause/title/risk/issue/suggestion/original。",
        params={"type": "object", "properties": {"contractName": {"type": "string"}}},
        handler=risk_scan_handler,
        risk="medium",
        skill="contract-review",
    ),
    ToolSpec(
        name="contract.redline_docx",
        title="导出带修订痕迹的合同",
        description="生成含 OOXML 修订痕迹（w:ins/w:del）与红色修订建议的 docx，需 HITL 审批。",
        params={"type": "object", "properties": {}},
        handler=redline_docx_handler,
        risk="high",
        requires_approval=True,
        skill="contract-review",
    ),
]
