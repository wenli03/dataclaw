"""数据源服务：路径解析、连通性测试、表/字段元数据（按权限过滤）。"""
from __future__ import annotations

import sqlite3
import time
from pathlib import Path
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..core.config import DATA_DIR, REPO
from ..models import ColumnMeta, DataSource, Policy, TableMeta, User, new_id
from . import policy_service

SUPPORTED_TYPES = ("sqlite", "mysql", "pg", "hive", "excel", "bi")
EXCEL_SUFFIXES = (".xlsx", ".xlsm")


def resolve_path(ds: DataSource) -> Path:
    """把 config.path 解析成绝对路径（兼容相对仓库根 / 相对 data 目录 / 绝对路径）。"""
    cfg: dict = ds.config or {}
    raw = str(cfg.get("path") or "").strip()
    if not raw:
        raw = ds.database or ""
    if not raw:
        return Path("")
    p = Path(raw)
    if p.is_absolute():
        return p
    # 相对路径依次尝试：
    #   1) 数据目录下的同名文件（DATACLAW_DATA_DIR，seed 实际落盘位置）
    #   2) 仓库根 / 仓库 data 目录（config 里写 data/demo.db 的常规情况）
    #   3) 数据目录 / backend 根（兜底）
    name = Path(raw).name
    candidates = [
        DATA_DIR / name,
        REPO / raw,
        REPO / "data" / name,
        DATA_DIR / raw,
        REPO / "backend" / raw,
    ]
    for cand in candidates:
        if cand.exists():
            return cand.resolve()
    return (DATA_DIR / name).resolve()


def demo_datasource(db: Session) -> DataSource | None:
    # SQLite JSON 查询在跨版本上有差异，直接按创建顺序取第一条（seed 只会建一条 demo）
    return db.scalar(select(DataSource).order_by(DataSource.created_at))


def get(db: Session, datasource_id: str | None = None) -> DataSource | None:
    if datasource_id:
        ds = db.get(DataSource, datasource_id)
        if ds:
            return ds
    return db.scalar(select(DataSource).order_by(DataSource.created_at))


def _table_count(db: Session, ds_id: str) -> int:
    return int(db.scalar(select(func.count()).select_from(TableMeta).where(TableMeta.datasource_id == ds_id)) or 0)


def to_dict(ds: DataSource, db: Session | None = None) -> dict:
    cfg = dict(ds.config or {})
    out = {
        "id": ds.id,
        "name": ds.name,
        "type": ds.type,
        "host": ds.host or "",
        "port": ds.port,
        "database": ds.database or "",
        "username": ds.username or "",
        "status": ds.status,
        "ownerId": ds.owner_id,
        "description": ds.description or "",
        "config": cfg,
        "path": str(resolve_path(ds)),
        "sourcePath": str(cfg.get("sourcePath") or ""),
        "createdAt": ds.created_at.isoformat() + "Z" if ds.created_at else None,
        "updatedAt": ds.updated_at.isoformat() + "Z" if ds.updated_at else None,
    }
    if db is not None:
        out["tableCount"] = _table_count(db, ds.id)
        out["rowCount"] = int(
            db.scalar(select(func.coalesce(func.sum(TableMeta.row_count), 0)).where(TableMeta.datasource_id == ds.id)) or 0
        )
        out["columnCount"] = int(
            db.scalar(select(func.count()).select_from(ColumnMeta).where(ColumnMeta.datasource_id == ds.id)) or 0
        )
        out["sensitiveColumnCount"] = int(
            db.scalar(
                select(func.count())
                .select_from(ColumnMeta)
                .where(
                    ColumnMeta.datasource_id == ds.id,
                    ColumnMeta.sensitive_level.in_(("secret", "top_secret")),
                )
            ) or 0
        )
    return out


def list_datasources(db: Session, user: User) -> dict:
    rows = db.scalars(select(DataSource).order_by(DataSource.created_at)).all()
    items = []
    for ds in rows:
        item = to_dict(ds, db)
        # 表可见数（按当前用户权限）
        try:
            visible = [t for t in policy_service.resolve_all(db, user, ds.id) if t.allowed]
            item["visibleTableCount"] = len(visible)
        except Exception:
            item["visibleTableCount"] = 0
        items.append(item)
    return {"items": items, "total": len(items)}


# ---------------------------------------------------------------- 表 / 字段


def visible_tables(db: Session, user: User, ds: DataSource) -> dict:
    """只返回当前用户有权限的表；同时给出列数与敏感列数。"""
    metas = db.scalars(
        select(TableMeta).where(TableMeta.datasource_id == ds.id).order_by(TableMeta.name)
    ).all()
    resolutions = {r.table: r for r in policy_service.resolve_all(db, user, ds.id)}
    items: list[dict] = []
    denied = 0
    for meta in metas:
        res = resolutions.get(meta.name)
        if res is None or not res.allowed:
            denied += 1
            continue
        items.append(
            {
                "name": meta.name,
                "comment": meta.comment or "",
                "rowCount": int(meta.row_count or 0),
                "columnCount": len(res.columns),
                "visibleColumnCount": len(res.visible_columns),
                "maskedColumnCount": len(res.masked_columns),
                "deniedColumnCount": len(res.denied_columns),
                "sensitiveColumns": [
                    c for c, rule in res.columns.items()
                    if rule.sensitive_level in ("secret", "top_secret")
                ],
                "rowFilters": res.row_filters,
                "allowed": True,
            }
        )
    return {
        "items": items,
        "total": len(items),
        "datasourceId": ds.id,
        "deniedColumnCount": 0,
        "hiddenTableCount": denied,
    }


def columns_of(db: Session, user: User, ds: DataSource, table: str) -> dict:
    """字段元数据（含敏感级别与当前用户的有效动作/脱敏方式）。"""
    res = policy_service.resolve(db, user, ds.id, table)
    if not res.allowed:
        raise PermissionError(res.denied_reason or f"无权访问表 {table}")
    metas = db.scalars(
        select(ColumnMeta)
        .where(ColumnMeta.datasource_id == ds.id, ColumnMeta.table_name == table)
        .order_by(ColumnMeta.name)
    ).all()
    items: list[dict] = []
    for cm in metas:
        rule = res.column_rule(cm.name)
        items.append(
            {
                "name": cm.name,
                "type": cm.type,
                "comment": cm.comment or "",
                "sensitiveLevel": cm.sensitive_level or "public",
                "category": cm.category or "",
                "action": (rule.action if rule else "allow"),
                "mask": (rule.mask if rule else "none"),
                "maskParam": dict(rule.mask_param or {}) if rule else {},
                "visible": bool(rule.visible) if rule else True,
                "reason": rule.reason if rule else "默认可见",
            }
        )
    return {
        "items": items,
        "total": len(items),
        "table": table,
        "datasourceId": ds.id,
        "rowFilters": res.row_filters,
        "allowed": True,
    }


# ---------------------------------------------------------------- Excel

def _quote_ident(name: str) -> str:
    """SQLite 标识符转义（工作表名允许中文与空格）。"""
    return '"' + str(name).replace('"', '""') + '"'


def _safe_ident(name: str, fallback: str) -> str:
    raw = str(name or "").strip()
    if not raw:
        return fallback
    # 保留中英文/数字/下划线，其余（空格、括号、斜杠）转下划线
    cleaned = "".join(ch if (ch.isalnum() or ch == "_") else "_" for ch in raw)
    cleaned = cleaned.strip("_")
    return cleaned or fallback


def _excel_value(v: Any) -> Any:
    from datetime import date, datetime, time

    if v is None or isinstance(v, (int, float, str)):
        return v
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (datetime, date, time)):
        return v.isoformat()
    return str(v)


def _infer_type(values: list[Any]) -> str:
    sample = [v for v in values if v is not None and str(v).strip() != ""][:80]
    if not sample:
        return "TEXT"
    if all(isinstance(v, bool) for v in sample):
        return "INTEGER"
    if all(isinstance(v, int) and not isinstance(v, bool) for v in sample):
        return "INTEGER"
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in sample):
        return "REAL"
    return "TEXT"


def _classify_column(name: str) -> tuple[str, str]:
    """按列名做轻量敏感识别：返回 (密级, 分类)。管理员仍可在策略里覆盖。"""
    n = str(name or "").lower()
    if any(k in n for k in ("身份证", "银行卡", "密码", "私钥", "社保")):
        return "top_secret", "个人信息"
    if any(k in n for k in ("薪资", "工资", "薪酬", "手机", "电话", "邮箱", "住址", "联系人")):
        return "secret", "个人信息"
    if any(k in n for k in ("金额", "成本", "利润", "营收", "合同额", "单价", "报价")):
        return "internal", "财务"
    if any(k in n for k in ("客户", "员工", "姓名", "负责人", "部门", "区域")):
        return "internal", "业务"
    return "public", "业务"


def import_excel(source: Path, target: Path) -> list[dict]:
    """把 Excel 各工作表导入为只读 SQLite 表，返回工作表摘要。

    DataClaw 的网关只接受 SQLite 连接；Excel 作为「接入格式」先落成可审计的
    本地只读库，后续表/列/行级策略、SQL 改写与脱敏全部复用同一套网关。
    """
    from openpyxl import load_workbook

    wb = load_workbook(str(source), read_only=True, data_only=True)
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        target.unlink()
    conn = sqlite3.connect(str(target))
    sheets: list[dict] = []
    used_tables: set[str] = set()
    try:
        for idx, ws in enumerate(wb.worksheets, start=1):
            rows = list(ws.iter_rows(values_only=True))
            while rows and all(c is None or str(c).strip() == "" for c in rows[0]):
                rows.pop(0)
            if not rows:
                continue
            header = list(rows[0])
            body = rows[1:]
            # 表头补全 / 去重
            columns: list[str] = []
            seen: dict[str, int] = {}
            for i, h in enumerate(header, start=1):
                col = _safe_ident(h, f"列{i}")
                seen[col] = seen.get(col, 0) + 1
                if seen[col] > 1:
                    col = f"{col}_{seen[col]}"
                columns.append(col)
            if not columns:
                continue
            table = _safe_ident(ws.title, f"Sheet{idx}")
            base = table
            n = 2
            while table in used_tables:
                table = f"{base}_{n}"
                n += 1
            used_tables.add(table)

            types = [_infer_type([row[i] if i < len(row) else None for row in body]) for i in range(len(columns))]
            ddl = ", ".join(f"{_quote_ident(c)} {t}" for c, t in zip(columns, types))
            conn.execute(f"CREATE TABLE {_quote_ident(table)} ({ddl})")
            placeholders = ", ".join("?" for _ in columns)
            insert_sql = f"INSERT INTO {_quote_ident(table)} VALUES ({placeholders})"
            payload = []
            for row in body:
                vals = [_excel_value(row[i] if i < len(row) else None) for i in range(len(columns))]
                if all(v is None or str(v).strip() == "" for v in vals):
                    continue
                payload.append(vals)
            if payload:
                conn.executemany(insert_sql, payload)
            sheets.append({
                "table": table,
                "sheet": ws.title,
                "columns": len(columns),
                "rows": len(payload),
                "columnDefs": [{"name": c, "type": t} for c, t in zip(columns, types)],
            })
        conn.commit()
    finally:
        conn.close()
        wb.close()
    return sheets


def _resolve_source_path(raw: str) -> Path:
    p = Path(raw)
    if p.is_absolute():
        return p
    for cand in (DATA_DIR / raw, REPO / raw, REPO / "data" / raw, Path.cwd() / raw):
        if cand.exists():
            return cand.resolve()
    return (DATA_DIR / Path(raw).name).resolve()


# ---------------------------------------------------------------- 连通性


def test_connection(db: Session, ds: DataSource) -> dict:
    started = time.perf_counter()
    result: dict[str, Any] = {
        "datasourceId": ds.id,
        "type": ds.type,
        "ok": False,
        "latencyMs": 0,
        "message": "",
    }
    try:
        if ds.type in ("sqlite", "excel"):
            path = resolve_path(ds)
            if not str(path) or not path.exists():
                result["message"] = f"SQLite 文件不存在：{path}"
            else:
                conn = sqlite3.connect(f"file:{path.as_posix()}?mode=ro", uri=True, timeout=5.0)
                try:
                    row = conn.execute("select 1").fetchone()
                    tables = conn.execute(
                        "select count(*) from sqlite_master where type='table'"
                    ).fetchone()[0]
                finally:
                    conn.close()
                label = "Excel 导入库" if ds.type == "excel" else "SQLite"
                result.update(
                    ok=bool(row and row[0] == 1),
                    message=f"{label}连接成功，共 {tables} 张表",
                    tableCount=int(tables),
                    path=str(path),
                )
        elif ds.type in ("mysql", "pg", "hive"):
            result["message"] = (
                f"V1 仅演示 sqlite：{ds.type} 驱动未随原型内置，"
                "请在 requirements 中安装对应驱动后启用（连接串已保存）"
            )
            result["retryable"] = False
        else:
            result["message"] = f"暂不支持的数据源类型：{ds.type}"
    except Exception as exc:
        result["message"] = f"连接失败：{exc}"
        result["retryable"] = True

    result["latencyMs"] = int((time.perf_counter() - started) * 1000)
    ds.status = "connected" if result["ok"] else "error"
    db.commit()
    return result


# ---------------------------------------------------------------- 创建


def create(db: Session, payload: dict, user: User) -> DataSource:
    ds_type = str(payload.get("type") or "sqlite").lower()
    if ds_type not in SUPPORTED_TYPES:
        raise ValueError(f"不支持的数据源类型：{ds_type}（可选：{'/'.join(SUPPORTED_TYPES)}）")
    name = str(payload.get("name") or "").strip()
    if not name:
        raise ValueError("数据源名称不能为空")
    config = dict(payload.get("config") or {})
    sheets: list[dict] = []
    if ds_type == "sqlite":
        path = str(config.get("path") or payload.get("database") or "").strip()
        if not path:
            raise ValueError("SQLite 数据源必须提供文件路径（config.path）")
        if not Path(path).is_absolute():
            config["path"] = path.replace("\\", "/")
        else:
            config["path"] = path
    elif ds_type == "excel":
        raw = str(config.get("path") or payload.get("database") or "").strip()
        if not raw:
            raise ValueError("Excel 数据源必须提供 .xlsx 文件路径（库名 / 路径）")
        source = _resolve_source_path(raw)
        if not source.exists():
            raise ValueError(f"Excel 文件不存在：{source}")
        if source.suffix.lower() not in EXCEL_SUFFIXES:
            raise ValueError("Excel 数据源仅支持 .xlsx / .xlsm（.xls 请另存为 .xlsx）")
        ds_id = new_id("ds")
        target = DATA_DIR / f"excel_{ds_id}.db"
        try:
            sheets = import_excel(source, target)
        except Exception:
            target.unlink(missing_ok=True)
            raise
        if not sheets:
            target.unlink(missing_ok=True)
            raise ValueError("未在 Excel 中找到可导入的工作表（请确认首行为表头且至少有一行数据）")
        config = {
            "path": str(target),
            "sourcePath": str(source),
            "readonly": True,
            "imported": True,
            "sheets": sheets,
        }
    else:
        ds_id = new_id("ds")
    ds = DataSource(
        id=ds_id,
        name=name,
        type=ds_type,
        host=str(payload.get("host") or ""),
        port=int(payload["port"]) if payload.get("port") else None,
        database=str(payload.get("database") or ""),
        username=str(payload.get("username") or ""),
        config=config,
        status="connected" if ds_type in ("sqlite", "excel") else "error",
        owner_id=user.id,
        description=str(payload.get("description") or ""),
    )
    db.add(ds)
    db.flush()
    if ds_type == "excel":
        for sheet in sheets:
            db.add(
                TableMeta(
                    datasource_id=ds.id,
                    name=sheet["table"],
                    comment=f"Excel 工作表：{sheet['sheet']}",
                    row_count=int(sheet["rows"]),
                )
            )
            for col in sheet.get("columnDefs") or []:
                level, category = _classify_column(col["name"])
                db.add(
                    ColumnMeta(
                        datasource_id=ds.id,
                        table_name=sheet["table"],
                        name=col["name"],
                        type=col["type"],
                        comment=f"来自 {Path(config['sourcePath']).name} / {sheet['sheet']}",
                        sensitive_level=level,
                        category=category,
                    )
                )
        # 新数据源默认只有平台管理员可见，业务用户需显式配置 allow 策略。
        db.add(
            Policy(
                id=new_id("pol"),
                subject_type="role",
                subject_id="platform_admin",
                resource_type="datasource",
                resource_id=ds.id,
                action="allow",
                priority=5,
                enabled=True,
                note="平台管理员：新接入数据源可见（业务用户需单独授权）",
            )
        )
    db.commit()
    db.refresh(ds)
    return ds



__all__ = [
    "resolve_path", "demo_datasource", "get", "to_dict", "list_datasources",
    "visible_tables", "columns_of", "test_connection", "create", "import_excel",
    "SUPPORTED_TYPES", "EXCEL_SUFFIXES",
]
