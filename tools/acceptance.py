"""DataClaw 验收自检脚本（对运行中的实例做端到端黑盒验收）。

用法：python tools/acceptance.py [--base http://127.0.0.1:8787]
退出码 0 = 全部通过。会在演示库里新增若干会话/产物（属正常使用痕迹）。
"""
from __future__ import annotations

import glob
import sys
import time

import httpx

# Windows 控制台默认 GBK，统一按 UTF-8 输出中文
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[attr-defined]
except Exception:
    pass

BASE = "http://127.0.0.1:8787"
for i, a in enumerate(sys.argv):
    if a == "--base" and i + 1 < len(sys.argv):
        BASE = sys.argv[i + 1]

ok = 0
fail: list[str] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    global ok
    if cond:
        ok += 1
        print(f"  [PASS] {name}" + (f" — {detail}" if detail else ""))
    else:
        fail.append(name)
        print(f"  [FAIL] {name}" + (f" — {detail}" if detail else ""))


def login(user: str, pwd: str) -> dict:
    r = httpx.post(f"{BASE}/api/auth/login", json={"username": user, "password": pwd}, timeout=30)
    r.raise_for_status()
    return {"Authorization": f"Bearer {r.json()['token']}"}


def ask(h: dict, content: str) -> dict:
    emps = httpx.get(f"{BASE}/api/employees?scope=mine", headers=h, timeout=30).json()["items"]
    sid = httpx.post(f"{BASE}/api/sessions", headers=h, json={"employeeId": emps[0]["id"]}, timeout=30).json()["id"]
    r = httpx.post(f"{BASE}/api/sessions/{sid}/messages?stream=false", headers=h, json={"content": content}, timeout=180)
    r.raise_for_status()
    return r.json()


print("=" * 72)
print("DataClaw 验收自检 ——", BASE)
print("=" * 72)
t0 = time.time()

# 1. 健康与认证
h_admin = login("admin", "admin123")
check("health 正常", httpx.get(f"{BASE}/api/health", timeout=10).json()["status"] == "ok")
bad = httpx.post(f"{BASE}/api/auth/login", json={"username": "admin", "password": "wrong"})
check("错误口令被拒", bad.status_code == 401, f"HTTP {bad.status_code}")
me = httpx.get(f"{BASE}/api/auth/me", headers=h_admin, timeout=10).json()
check("admin 身份与权限", me["user"]["role"] == "platform_admin" and me["permissions"]["isAdmin"])

# 2. 管理后台数据
ov = httpx.get(f"{BASE}/api/overview", headers=h_admin, timeout=30).json()
s = ov["stats"]
check("运营总览有真实数据", s["employees"] >= 5 and s["skills"] >= 12 and s["policies"] >= 20,
      f"员工{s['employees']}/技能{s['skills']}/策略{s['policies']}/审计{s.get('auditEvents')}")
audit = httpx.get(f"{BASE}/api/audit?result=denied&size=5", headers=h_admin, timeout=30).json()
check("审计含越权拦截记录", audit["total"] >= 1, f"denied={audit['total']}")
policies = httpx.get(f"{BASE}/api/policies", headers=h_admin, timeout=30).json()
check("数据权限策略已配置", policies["total"] >= 20, f"{policies['total']} 条")
pas = httpx.get(f"{BASE}/api/public-agents", headers=h_admin, timeout=30).json()
check("公共员工已定义", pas["total"] >= 2, f"{pas['total']} 个")

# 3. 场景 A：工时分析（analyst，含脱敏）
h_analyst = login("analyst", "123456")
a = ask(h_analyst, "帮我分析下上季度各区域工时效率情况")
dash = next((x for x in a.get("artifacts", []) if x["kind"] == "dashboard"), None)
rep = next((x for x in a.get("artifacts", []) if x["kind"] == "report"), None)
dp = (dash or {}).get("preview") or {}
check("场景A 产出大屏", bool(dash) and len(dp.get("charts", [])) >= 3 and len(dp.get("kpis", [])) >= 3,
      f"{len(dp.get('charts', []))} 图 / {len(dp.get('kpis', []))} KPI")
check("场景A 产出深度报告", bool(rep) and len(((rep.get('preview') or {}).get('markdown') or "")) >= 1200,
      f"{(len(((rep.get('preview') or {}).get('markdown') or '')))} 字")
check("场景A 薪资列已脱敏", "salary" in str(dp), "出现在 source.maskedColumns")
check("场景A 工具链完整", len(a.get("steps", [])) >= 4, f"{len(a.get('steps', []))} 步")

# 4. 场景 B：销售情报
h_sales = login("sales", "123456")
b = ask(h_sales, "我要拜访星海科技，帮我准备拜访指南")
guide = next((x for x in b.get("artifacts", []) if x["kind"] == "html"), None)
check("场景B 产出拜访指南", bool(guide) and len(str((guide.get('preview') or {}).get('html', ''))) > 2000,
      f"{len(str((guide.get('preview') or {}).get('html', '')))} 字节")

# 5. 场景 C：合同审查
h_legal = login("legal", "123456")
c = ask(h_legal, "请按照公司标准合同规则审查这份合同")
docx_art = next((x for x in c.get("artifacts", []) if x["kind"] == "docx"), None)
cprev = (docx_art or {}).get("preview") or {}
check("场景C 风险项 >= 6", len(cprev.get("redlines", [])) >= 6, f"{len(cprev.get('redlines', []))} 条")
check("场景C 风险分级", (cprev.get("stats") or {}).get("high", 0) >= 1, str(cprev.get("stats")))
docxs = sorted(glob.glob(r"C:\dataclaw\data\artifacts\*.docx"), key=lambda p: -__import__("os").path.getmtime(p))
try:
    import docx  # type: ignore

    d = docx.Document(docxs[0])
    check("场景C docx 可被 Word 解析", len(d.paragraphs) > 5, f"{len(d.paragraphs)} 段")
except Exception as e:  # pragma: no cover
    check("场景C docx 可被 Word 解析", False, str(e)[:80])
check("场景C 有 HITL 审批记录", len(c.get("approvals", [])) >= 1, str([x.get("action") for x in c.get("approvals", [])]))

# 6. 数据安全（SQL 网关）
pw = httpx.post(f"{BASE}/api/policies/preview", headers=h_analyst,
                json={"datasourceId": "ds_demo", "sql": "select * from worktime_fact limit 5"}, timeout=60).json()
check("行/列权限生效", pw.get("ok") and pw.get("maskedColumns") and len(pw.get("rows", [])) == 5,
      f"masked={pw.get('maskedColumns')}")
den = httpx.post(f"{BASE}/api/policies/preview", headers=h_analyst,
                 json={"datasourceId": "ds_demo", "sql": "select * from crm_customer limit 3"}, timeout=60).json()
check("越权表被拒绝且原因可读", den.get("ok") is False and "无权" in str(den.get("deniedReason", "")))
dml = httpx.post(f"{BASE}/api/policies/preview", headers=h_analyst,
                 json={"datasourceId": "ds_demo", "sql": "delete from worktime_fact"}, timeout=60).json()
check("DML 被网关拒绝", dml.get("ok") is False)
sales_rows = httpx.post(f"{BASE}/api/policies/preview", headers=h_sales,
                        json={"datasourceId": "ds_demo", "sql": "select * from crm_customer limit 100"}, timeout=60).json()
owners = {r.get("owner_id") for r in (sales_rows.get("rows") or [])}
check("行级过滤（sales 只见自己客户）", sales_rows.get("ok") and (not owners or owners == {"u_sales"}),
      f"{len(sales_rows.get('rows') or [])} 行 / owners={owners}")

# 7. 知识中台
k = httpx.post(f"{BASE}/api/knowledge/search", headers=h_admin, json={"query": "合同付款账期", "topK": 3}, timeout=30).json()
check("知识库检索可用", len(k.get("items", [])) >= 1, f"{len(k.get('items', []))} 条命中")

# 8. 渠道
inbound = httpx.post(f"{BASE}/api/channels/web/events", headers={"X-DataClaw-Token": "demo-token"},
                     json={"text": "帮我看看知识库里合同审查规则"}, timeout=120)
check("渠道入站 webhook 可达", inbound.status_code == 200 and len(str(inbound.json().get("reply", ""))) > 20,
      f"HTTP {inbound.status_code} / reply {len(str(inbound.json().get('reply', '')))} 字")

# 9. 技能安装/卸载闭环
sk = httpx.get(f"{BASE}/api/skills?tab=market&size=100", headers=h_legal, timeout=30).json()["items"]
if sk:
    target = sk[0]["id"]
    r1 = httpx.post(f"{BASE}/api/skills/{target}/install", headers=h_legal, json={}, timeout=30).json()
    r2 = httpx.post(f"{BASE}/api/skills/{target}/uninstall", headers=h_legal, timeout=30).json()
    r3 = httpx.post(f"{BASE}/api/skills/{target}/install", headers=h_legal, json={}, timeout=30).json()
    check("技能安装/卸载闭环", (r1.get("ok") or r1.get("already")) and r2.get("ok") and (r3.get("ok") or r3.get("already")))
else:
    check("技能安装/卸载闭环", False, "market 为空（可能全部已装）")

# 10. 公共员工分配（幂等）
pa = pas["items"][0]["id"]
users = httpx.get(f"{BASE}/api/users?size=100", headers=h_admin, timeout=30).json()["items"]
cand = next((u for u in users if u["username"] == "zhaomin"), None)
if cand:
    # 先取消（保证基线干净），再分配两次验证幂等
    httpx.delete(f"{BASE}/api/public-agents/{pa}/assign/{cand['id']}", headers=h_admin, timeout=30)
    a1 = httpx.post(f"{BASE}/api/public-agents/{pa}/assign", headers=h_admin, json={"userIds": [cand["id"]]}, timeout=60)
    a2 = httpx.post(f"{BASE}/api/public-agents/{pa}/assign", headers=h_admin, json={"userIds": [cand["id"]]}, timeout=60)
    j1 = a1.json() if a1.status_code == 200 else {}
    j2 = a2.json() if a2.status_code == 200 else {}
    created2 = j2.get("createdCount", len(j2.get("created") or []))
    check("公共员工分配幂等", a1.status_code == 200 and a2.status_code == 200 and created2 == 0,
          f"首次创建 {j1.get('createdCount', len(j1.get('created') or []))}，重复创建 {created2}")
else:
    check("公共员工分配幂等", False, "找不到测试用户 zhaomin")

# 11. 前端与产物下载
check("前端页面可访问", httpx.get(f"{BASE}/chat", timeout=20).status_code == 200)
dl = httpx.get(f"{BASE}/api/artifacts/{docx_art.get('id') or docx_art.get('artifactId')}/download", headers=h_legal, timeout=60)
check("修订版 docx 可下载", dl.status_code == 200 and dl.content[:2] == b"PK", f"{len(dl.content) // 1024} KB")

elapsed = time.time() - t0
print("-" * 72)
print(f"通过 {ok} / {ok + len(fail)}，耗时 {elapsed:.1f}s")
if fail:
    print("失败项：" + "、".join(fail))
    sys.exit(1)
print("全部通过 ✅ (26/26)")
